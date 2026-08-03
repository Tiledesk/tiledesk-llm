#!/usr/bin/env python3
"""
Converters: source bytes/text -> ExtractedDocument (canonical model).

txt/md/csv/xlsx are pure/sync (no heavy deps). pdf/docx delegate to the
existing engines (docling seam from pdf_ocr, StructuredDocxLoader) via
dependency injection — the test venv has neither `docling` nor `docx`
installed (see memory/pdf_ocr_native_md_converter_seam), so those converters
must lazy-import the real engine and accept an injected fake for testing.
"""
from unittest.mock import Mock

import pytest

from tilellm.modules.ingestion.export.converters import (
    convert_csv,
    convert_docx,
    convert_md,
    convert_pdf,
    convert_txt,
    convert_url,
    convert_xlsx,
)


class TestConvertTxt:
    def test_single_text_block(self):
        doc = convert_txt("Hello world")
        assert doc.type == "Text Document"
        assert len(doc.blocks) == 1
        assert doc.blocks[0].content == "Hello world"
        assert doc.blocks[0].block_type == "text"

    def test_resource_propagated(self):
        doc = convert_txt("hi", resource="https://x/y.txt")
        assert doc.resource == "https://x/y.txt"


class TestConvertMd:
    def test_plain_markdown_wrapped(self):
        doc = convert_md("# Title\n\nSome content")
        assert doc.type == "Markdown Document"
        assert len(doc.blocks) == 1
        assert "# Title" in doc.blocks[0].content

    def test_existing_frontmatter_is_parsed_not_rewrapped(self):
        raw = "---\ntype: PDF Document\ntitle: Capitolato\n---\n\nBody text\n"
        doc = convert_md(raw)
        assert doc.type == "PDF Document"
        assert doc.title == "Capitolato"
        assert doc.blocks[0].content.strip() == "Body text"


class TestConvertCsv:
    def test_produces_single_table_block(self):
        csv_bytes = b"name,age\nMario,30\nLuca,25\n"
        doc = convert_csv(csv_bytes)
        assert doc.type == "Tabular Document"
        assert len(doc.blocks) == 1
        assert doc.blocks[0].block_type == "table"
        md = doc.blocks[0].content
        assert "name" in md and "age" in md
        assert "Mario" in md and "30" in md

    def test_no_tabulate_dependency_needed(self):
        # regression guard: must not raise ImportError even though `tabulate`
        # is not installed in this test environment
        convert_csv(b"a,b\n1,2\n")


class TestConvertXlsx:
    def test_one_block_per_sheet(self):
        import io
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1.append(["col1", "col2"])
        ws1.append(["v1", "v2"])
        ws2 = wb.create_sheet("Sheet2")
        ws2.append(["x"])
        ws2.append(["y"])
        buf = io.BytesIO()
        wb.save(buf)

        doc = convert_xlsx(buf.getvalue())
        assert doc.type == "Tabular Document"
        assert len(doc.blocks) == 2
        assert doc.blocks[0].heading_path == "Sheet1"
        assert "col1" in doc.blocks[0].content
        assert doc.blocks[1].heading_path == "Sheet2"


class TestConvertPdf:
    @pytest.mark.asyncio
    async def test_page_bodies_become_page_blocks(self):
        fake_result = Mock(page_bodies=[(1, "Pagina uno"), (2, "Pagina due")])
        fake_converter = Mock(side_effect=lambda *a, **k: _async_return(fake_result))

        doc = await convert_pdf("/tmp/fake.pdf", "doc1", converter=fake_converter)

        assert doc.type == "PDF Document"
        assert [b.content for b in doc.blocks] == ["Pagina uno", "Pagina due"]
        assert [b.page for b in doc.blocks] == [1, 2]
        assert all(b.block_type == "page" for b in doc.blocks)

    @pytest.mark.asyncio
    async def test_skip_ocr_option_forwarded(self):
        fake_result = Mock(page_bodies=[])
        calls = {}

        async def fake_converter(file_path, doc_id, attempt=1, options=None):
            calls["options"] = options
            return fake_result

        await convert_pdf("/tmp/f.pdf", "doc1", converter=fake_converter, skip_ocr=True)
        assert calls["options"] == {"skip_ocr": True}

    @pytest.mark.asyncio
    async def test_skip_ocr_none_classifies_native_as_ocr_off(self):
        """skip_ocr omitted (None) -> per-document classification decides, not a
        fixed default (docs/MIGLIORIE_DA_FARE.md UPGRADE trap: a fixed True would
        silently drop scanned content on non-native PDFs)."""
        fake_result = Mock(page_bodies=[])
        calls = {}

        async def fake_converter(file_path, doc_id, attempt=1, options=None):
            calls["options"] = options
            return fake_result

        fake_classifier = Mock(return_value={"doc_type": "native"})
        await convert_pdf("/tmp/f.pdf", "doc1", converter=fake_converter, classifier=fake_classifier)

        fake_classifier.assert_called_once_with("/tmp/f.pdf")
        assert calls["options"] == {"skip_ocr": True}

    @pytest.mark.asyncio
    async def test_skip_ocr_none_classifies_scanned_as_ocr_on(self):
        fake_result = Mock(page_bodies=[])
        calls = {}

        async def fake_converter(file_path, doc_id, attempt=1, options=None):
            calls["options"] = options
            return fake_result

        fake_classifier = Mock(return_value={"doc_type": "scanned"})
        await convert_pdf("/tmp/f.pdf", "doc1", converter=fake_converter, classifier=fake_classifier)

        assert calls["options"] == {"skip_ocr": False}

    @pytest.mark.asyncio
    async def test_classifier_failure_defaults_to_ocr_on(self):
        """Unknown classification -> never assume native; OCR stays on."""
        fake_result = Mock(page_bodies=[])
        calls = {}

        async def fake_converter(file_path, doc_id, attempt=1, options=None):
            calls["options"] = options
            return fake_result

        def broken_classifier(path):
            raise RuntimeError("boom")

        await convert_pdf("/tmp/f.pdf", "doc1", converter=fake_converter, classifier=broken_classifier)

        assert calls["options"] == {"skip_ocr": False}


class TestConvertUrl:
    @pytest.mark.asyncio
    async def test_documents_become_text_blocks(self):
        from langchain_core.documents import Document

        fake_docs = [Document(page_content="Pagina scrapata")]

        async def fake_fetch(source, scrape_type, parameters_scrape_type_4=None, browser_headers=None):
            return fake_docs

        doc = await convert_url("https://example.com", fetch=fake_fetch)

        assert doc.type == "Web Page"
        assert doc.resource == "https://example.com"
        assert [b.content for b in doc.blocks] == ["Pagina scrapata"]
        assert doc.blocks[0].block_type == "text"

    @pytest.mark.asyncio
    async def test_scrape_params_forwarded(self):
        calls = {}

        async def fake_fetch(source, scrape_type, parameters_scrape_type_4=None, browser_headers=None):
            calls["scrape_type"] = scrape_type
            calls["browser_headers"] = browser_headers
            return []

        await convert_url("https://x", scrape_type=2, browser_headers={"h": "v"}, fetch=fake_fetch)

        assert calls["scrape_type"] == 2
        assert calls["browser_headers"] == {"h": "v"}


class TestConvertDocx:
    def test_text_and_table_blocks(self):
        from langchain_core.documents import Document

        fake_docs = [
            Document(page_content="Paragrafo 1", metadata={"element_type": "text"}),
            Document(page_content="| a | b |", metadata={"element_type": "table"}),
        ]
        fake_loader_cls = Mock(return_value=Mock(load_with_images=Mock(return_value=(fake_docs, []))))

        doc = convert_docx("/tmp/f.docx", loader_cls=fake_loader_cls)

        assert doc.type == "Word Document"
        assert [b.block_type for b in doc.blocks] == ["text", "table"]
        assert doc.blocks[1].content == "| a | b |"

    def test_no_fake_page_number(self):
        """DOCX has no real pagination — Block.page must stay None, never guessed."""
        from langchain_core.documents import Document

        fake_docs = [Document(page_content="Testo", metadata={"element_type": "text", "_para_index": 4})]
        fake_loader_cls = Mock(return_value=Mock(load_with_images=Mock(return_value=(fake_docs, []))))

        doc = convert_docx("/tmp/f.docx", loader_cls=fake_loader_cls)
        assert doc.blocks[0].page is None

    def test_heading_path_and_para_index_propagated(self):
        from langchain_core.documents import Document

        fake_docs = [
            Document(
                page_content="Paragrafo",
                metadata={"element_type": "text", "heading_path": "Sezione 2", "_para_index": 7},
            ),
        ]
        fake_loader_cls = Mock(return_value=Mock(load_with_images=Mock(return_value=(fake_docs, []))))

        doc = convert_docx("/tmp/f.docx", loader_cls=fake_loader_cls)
        assert doc.blocks[0].heading_path == "Sezione 2"
        assert doc.blocks[0].position == 7

    def test_table_index_used_as_position_for_tables(self):
        from langchain_core.documents import Document

        fake_docs = [
            Document(
                page_content="| a |",
                metadata={"element_type": "table", "heading_path": "Sez", "table_index": 2},
            ),
        ]
        fake_loader_cls = Mock(return_value=Mock(load_with_images=Mock(return_value=(fake_docs, []))))

        doc = convert_docx("/tmp/f.docx", loader_cls=fake_loader_cls)
        assert doc.blocks[0].position == 2


async def _async_return(value):
    return value
