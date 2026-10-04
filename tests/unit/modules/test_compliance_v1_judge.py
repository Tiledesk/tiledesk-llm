"""
v1 Conformità judge: citation attribution and prompt contract.

Found on a real tender: a requirement judged "not found" was reported with a
document/page that had nothing to do with it — the judge anchored no chunk, and
_pick_best_source's last-resort fallback (first retrieved chunk) was written out as
if it were the evidence. The domain system prompts also told the model to answer
with "exactly" four keys, contradicting the user prompt that asks for
source_chunk_index as well.
"""
import io
import json
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest

from tilellm.modules.compliance_checker.logic import _JUDGE_USER_PROMPT, _judge_requirement
from tilellm.modules.compliance_checker.models import ComplianceResult, RequirementItem
from tilellm.modules.compliance_checker.prompts import (
    DISCRETIONARY_JUDGE_SYSTEM_PROMPT,
    get_builtin_config,
    list_builtin_domains,
)

_CHUNKS = ["Pagina di copertina e marchi.", "Il dispositivo è conforme al requisito richiesto."]
_METADATA = [
    {"file_name": "brochure.pdf", "page": 6},
    {"file_name": "scheda_tecnica.pdf", "page": 14},
]


def _llm_answering(payload: dict):
    response = MagicMock()
    response.content = json.dumps(payload)
    llm = AsyncMock()
    llm.ainvoke = AsyncMock(return_value=response)
    return llm


async def _judge(payload: dict) -> ComplianceResult:
    return await _judge_requirement(
        RequirementItem(id="C9", text="Requisito di prova"),
        _CHUNKS, _METADATA, get_builtin_config("e_procurement"), _llm_answering(payload),
    )


# ---------------------------------------------------------------------------
# Citation attribution
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_anchored_citation_points_to_the_cited_chunk():
    result = await _judge({
        "judgment": "compliant", "confidence": 0.9, "source_chunk_index": 2,
        "evidence_text": "conforme al requisito", "justification": "ok",
    })

    assert result.citation_attributed is True
    assert result.evidence_document == "scheda_tecnica.pdf"
    assert result.evidence_page == 14


@pytest.mark.asyncio
async def test_unanchored_citation_reports_no_document_instead_of_a_fallback():
    """The first retrieved chunk is NOT evidence just because nothing else anchored."""
    result = await _judge({
        "judgment": "not_verifiable", "confidence": 0.1, "source_chunk_index": 0,
        "evidence_text": "", "justification": "nessuna evidenza",
    })

    assert result.citation_attributed is False
    assert result.evidence_document == ""


def test_restitution_xlsx_flags_unattributed_tabular_citation():
    from tilellm.modules.compliance_checker.models_v2 import (
        ComplianceReportV2, ComplianceSummaryV2, TenderInfo,
    )
    from tilellm.modules.compliance_checker.services.restituzione_xlsx_service import (
        RestituzioneXlsxService,
    )

    tabular = ComplianceResult(
        requirement_id="C9", requirement_text="Requisito di prova", category=None, mandatory=True,
        judgment="not_verifiable", confidence=0.1, evidence_text="", justification="nessuna evidenza",
        evidence_document="", evidence_page=1, evidence_section="", citation_attributed=False,
    )
    report = ComplianceReportV2(
        tender=TenderInfo(title="t", lot_id="L1", lot_name="Lotto 1"), namespace="OE1",
        summary=ComplianceSummaryV2(), tabular_results=[tabular], discretionary_results=[],
    )

    workbook = openpyxl.load_workbook(io.BytesIO(RestituzioneXlsxService().build_workbook([report])))
    cells = [c.value for row in workbook.active.iter_rows() for c in row]

    assert "⚠ da verificare" in cells


# ---------------------------------------------------------------------------
# Prompt contract
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("domain", list_builtin_domains())
def test_domain_system_prompt_does_not_dictate_the_json_keys(domain):
    """The output format lives in ONE place (_JUDGE_USER_PROMPT). A second, shorter
    key list in the system prompt contradicted it and dropped source_chunk_index."""
    prompt = get_builtin_config(domain).system_prompt.lower()

    assert "exactly these keys" not in prompt
    assert "esattamente queste chiavi" not in prompt


@pytest.mark.parametrize("prompt", [_JUDGE_USER_PROMPT, DISCRETIONARY_JUDGE_SYSTEM_PROMPT],
                         ids=["v1_conformita", "v2_discrezionale"])
def test_judge_prompts_accept_equivalent_wording_and_other_languages(prompt):
    text = prompt.lower()

    assert "sinonim" in text
    assert "altra lingua" in text
    assert "dichiarazione" in text


@pytest.mark.parametrize("prompt", [_JUDGE_USER_PROMPT, DISCRETIONARY_JUDGE_SYSTEM_PROMPT],
                         ids=["v1_conformita", "v2_discrezionale"])
def test_judge_prompts_stay_domain_agnostic(prompt):
    """The equivalence rule must be generic — the service evaluates tenders of any
    sector, so no tender- or product-specific vocabulary may leak into it."""
    text = prompt.lower()

    for term in ("iso 10993", "biocompat", "cemento", "mdr"):
        assert term not in text


# ---------------------------------------------------------------------------
# Judge robustness — shared with the v2 judge
#
# Found on a real run: a provider answered one Conformità requirement with
# something json.loads rejected ("Expecting value"); v1 made a single attempt and
# the failure came out as a plain not_verifiable, indistinguishable from a real
# "no evidence" verdict. Credit exhaustion would have looked the same.
# ---------------------------------------------------------------------------

from tilellm.modules.compliance_checker import logic as _logic  # noqa: E402


@pytest.fixture(autouse=False)
def _no_retry_delay(monkeypatch):
    monkeypatch.setattr(_logic, "_JUDGE_RETRY_DELAY_S", 0)


def _llm_answering_raw(*contents):
    """LLM whose successive ainvoke calls return the given contents; an Exception
    instance is raised instead of returned."""
    side_effects = []
    for c in contents:
        if isinstance(c, Exception):
            side_effects.append(c)
        else:
            response = MagicMock()
            response.content = c
            side_effects.append(response)
    llm = AsyncMock()
    llm.ainvoke = AsyncMock(side_effect=side_effects)
    return llm


async def _judge_with(llm) -> ComplianceResult:
    return await _judge_requirement(
        RequirementItem(id="C9", text="Requisito di prova"),
        _CHUNKS, _METADATA, get_builtin_config("e_procurement"), llm,
    )


_COMPLIANT = {"judgment": "compliant", "confidence": 0.9, "source_chunk_index": 2,
              "evidence_text": "conforme al requisito", "justification": "ok"}


@pytest.mark.parametrize("raw", [
    "Ecco la valutazione:\n```json\n" + json.dumps(_COMPLIANT) + "\n```",
    json.dumps(_COMPLIANT) + "\n\nNota: {vedi sopra}",
], ids=["preamble_and_fence", "trailing_prose_with_braces"])
def test_parse_judge_json_finds_the_object_wherever_it_is(raw):
    assert _logic._parse_judge_json(raw) == _COMPLIANT


@pytest.mark.parametrize("raw", ["", "nessun json qui", "[1, 2]"])
def test_parse_judge_json_rejects_non_objects(raw):
    with pytest.raises(_logic.JudgeInvocationError):
        _logic._parse_judge_json(raw)


@pytest.mark.asyncio
async def test_v1_judge_retries_a_transient_failure(_no_retry_delay):
    llm = _llm_answering_raw(RuntimeError("overloaded"), json.dumps(_COMPLIANT))

    result = await _judge_with(llm)

    assert result.judgment == "compliant"
    assert result.judge_failed is False
    assert llm.ainvoke.await_count == 2


@pytest.mark.asyncio
async def test_v1_judge_failure_is_flagged_not_disguised_as_no_evidence(_no_retry_delay):
    llm = _llm_answering_raw(*(["non è json"] * _logic._MAX_JUDGE_ATTEMPTS))

    result = await _judge_with(llm)

    assert result.judge_failed is True
    assert result.judgment == "not_verifiable"
    assert "giudice" in result.justification.lower()
    assert result.evidence_document == ""


def test_restitution_xlsx_marks_a_failed_judge_for_human_review():
    from tilellm.modules.compliance_checker.models_v2 import (
        ComplianceReportV2, ComplianceSummaryV2, TenderInfo,
    )
    from tilellm.modules.compliance_checker.services.restituzione_xlsx_service import (
        RestituzioneXlsxService,
    )

    tabular = ComplianceResult(
        requirement_id="C9", requirement_text="Requisito di prova", category=None, mandatory=True,
        judgment="not_verifiable", confidence=0.0, evidence_text="",
        justification="Giudice LLM non disponibile", evidence_document="", evidence_page=1,
        evidence_section="", citation_attributed=False, judge_failed=True,
    )
    report = ComplianceReportV2(
        tender=TenderInfo(title="t", lot_id="L1", lot_name="Lotto 1"), namespace="OE1",
        summary=ComplianceSummaryV2(), tabular_results=[tabular], discretionary_results=[],
    )

    workbook = openpyxl.load_workbook(io.BytesIO(RestituzioneXlsxService().build_workbook([report])))
    cells = [c.value for row in workbook.active.iter_rows() for c in row]

    assert "⚠ REVISIONE UMANA" in cells


@pytest.mark.parametrize("prompt", [_JUDGE_USER_PROMPT, DISCRETIONARY_JUDGE_SYSTEM_PROMPT],
                         ids=["v1_conformita", "v2_discrezionale"])
def test_judge_prompts_are_strict_on_what_counts_as_evidence(prompt):
    """Tender decisions get appealed: after the equivalence rule the 28/09 run turned
    several confirmed PARZIALE/NO into SI — a repealed directive accepted for the
    regulation named in the requirement, a requirement with several elements passed
    on one, a characteristic inferred from a usage warning. Equivalence is about
    wording, never about substituting or inferring content."""
    text = " ".join(prompt.lower().split())

    assert "abrogat" in text                  # a different/repealed norm doesn't satisfy it
    assert "non dedurre" in text              # no inference from warnings, similar products...
    assert "ciascuno" in text                 # every listed element must be documented
    assert "nel dubbio" in text               # doubt -> not satisfied


@pytest.mark.parametrize("prompt", [_JUDGE_USER_PROMPT, DISCRETIONARY_JUDGE_SYSTEM_PROMPT],
                         ids=["v1_conformita", "v2_discrezionale"])
def test_judge_prompts_tell_cumulative_elements_from_alternatives(prompt):
    """Real run, 29/09: "conformi al Regolamento 745/2017 o alla Direttiva 93/42" was
    judged PARZIALE — "il requisito richiede conformità a entrambi" — although the
    offer documented 93/42. The every-element rule applies to elements required
    together, not to alternatives, where one documented option is enough."""
    text = " ".join(prompt.lower().split())

    assert "alternative" in text
    assert "ne basta una" in text
